from pathlib import Path
import os

import pytest
from tokenizers import Tokenizer, models, pre_tokenizers, trainers

from app.services.rag import documents
from app.services.rag.errors import KnowledgeBaseError


def parse(path: Path, media_type: str):
    function = getattr(documents, "parse_document", None)
    assert function is not None, "尚未实现文档解析"
    return function(path, media_type)


def split(blocks, tokenizer: Tokenizer, size: int = 128, overlap: int = 20):
    function = getattr(documents, "split_blocks", None)
    assert function is not None, "尚未实现 Token 切分"
    return function(blocks, tokenizer, size, overlap)


@pytest.fixture
def byte_tokenizer() -> Tokenizer:
    tokenizer = Tokenizer(models.BPE(unk_token="[UNK]"))
    tokenizer.pre_tokenizer = pre_tokenizers.ByteLevel(add_prefix_space=False)
    tokenizer.train_from_iterator(["中文 Alpha beta!"], trainers.BpeTrainer(
        vocab_size=257, special_tokens=["[UNK]"], initial_alphabet=pre_tokenizers.ByteLevel.alphabet(),
    ))
    return tokenizer


def test_text_parser_preserves_original_line_numbers_and_utf8(tmp_path: Path) -> None:
    path = tmp_path / "资料.txt"
    path.write_bytes("\ufeff第一行\r\n\r\n第三行，中文😀\n".encode())
    blocks = parse(path, "text/plain")
    assert len(blocks) == 1 and blocks[0].text == "第一行\n\n第三行，中文😀\n"
    assert (blocks[0].source.line_start, blocks[0].source.line_end) == (1, 3)


def test_markdown_splitting_preserves_content_locations_and_token_limit(tmp_path: Path, byte_tokenizer: Tokenizer) -> None:
    path = tmp_path / "资料.md"
    original = "# 制度\n\n" + "报销需发票，Approval required!😀\n" * 40 + "\n## 注意\n不得丢失尾部。"
    path.write_text(original, encoding="utf-8")
    chunks = split(parse(path, "text/markdown"), byte_tokenizer)
    assert len(chunks) > 2
    coverage = set()
    lines = original.splitlines()
    for chunk in chunks:
        assert 0 < chunk.token_count <= 128
        assert chunk.token_count == len(byte_tokenizer.encode(chunk.text, add_special_tokens=True).ids)
        assert "�" not in chunk.text
        source_text = "\n".join(lines[chunk.source.line_start - 1:chunk.source.line_end])
        assert chunk.text in source_text
        coverage.update(range(chunk.source.line_start, chunk.source.line_end + 1))
    assert {index for index, line in enumerate(lines, 1) if line.strip()} <= coverage
    assert "不得丢失尾部。" in chunks[-1].text


def test_long_unbroken_unicode_is_split_without_losing_characters(tmp_path: Path, byte_tokenizer: Tokenizer) -> None:
    path = tmp_path / "长句.txt"
    original = "中文🧑🏽‍💻abcdef" * 100
    path.write_text(original, encoding="utf-8")
    chunks = split(parse(path, "text/plain"), byte_tokenizer, overlap=0)
    assert "".join(chunk.text for chunk in chunks) == original
    assert all(chunk.token_count <= 128 for chunk in chunks)


def test_docx_parser_keeps_paragraph_table_order_and_nested_cell_content(tmp_path: Path) -> None:
    from docx import Document
    document = Document()
    document.add_heading("制度标题", 1)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "姓名"
    nested = table.cell(0, 1).add_table(rows=1, cols=1)
    nested.cell(0, 0).text = "嵌套值"
    document.add_paragraph("最后说明")
    path = tmp_path / "资料.docx"
    document.save(path)
    blocks = parse(path, documents.MEDIA_TYPES[".docx"])
    assert len(blocks) == 3
    assert blocks[0].text == "制度标题"
    assert "姓名" in blocks[1].text and "嵌套值" in blocks[1].text
    assert blocks[2].text == "最后说明"
    assert [(item.source.block_start, item.source.block_type) for item in blocks] == [(1, "paragraph"), (2, "table"), (3, "paragraph")]


def make_pdf(path: Path, *, blank: bool = False, encrypted: bool = False) -> None:
    from pypdf import PdfWriter
    from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject
    writer = PdfWriter()
    for text in ["First page", "Second page"]:
        page = writer.add_blank_page(width=300, height=300)
        if not blank:
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"), NameObject("/Subtype"): NameObject("/Type1"), NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 10 250 Td ({text}) Tj ET".encode())
            page[NameObject("/Contents")] = stream
    if encrypted:
        writer.encrypt("test-only")
    writer.write(path)


def test_pdf_page_sources_and_cross_page_chunk(tmp_path: Path, byte_tokenizer: Tokenizer) -> None:
    path = tmp_path / "pages.pdf"
    make_pdf(path)
    blocks = parse(path, "application/pdf")
    assert [item.text for item in blocks] == ["First page", "Second page"]
    assert [item.source.page_start for item in blocks] == [1, 2]
    chunks = split(blocks, byte_tokenizer)
    assert len(chunks) == 1 and chunks[0].source.page_start == 1 and chunks[0].source.page_end == 2


@pytest.mark.parametrize(("blank", "encrypted", "code"), [(True, False, "pdf_no_text"), (False, True, "encrypted_document")])
def test_pdf_rejects_scanned_or_encrypted_content(tmp_path: Path, blank: bool, encrypted: bool, code: str) -> None:
    path = tmp_path / "pages.pdf"
    make_pdf(path, blank=blank, encrypted=encrypted)
    with pytest.raises(KnowledgeBaseError) as error:
        parse(path, "application/pdf")
    assert error.value.code == code


@pytest.mark.parametrize("extension", [".pdf", ".docx"])
def test_invalid_document_has_sanitized_error(tmp_path: Path, extension: str) -> None:
    path = tmp_path / ("内部文件名" + extension)
    path.write_bytes(b"invalid")
    with pytest.raises(KnowledgeBaseError) as error:
        parse(path, documents.MEDIA_TYPES[extension])
    assert error.value.code == "invalid_file" and str(path) not in str(error.value)


def test_parser_does_not_silently_truncate_extracted_content(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    path = tmp_path / "large.txt"
    path.write_text("a" * 500, encoding="utf-8")
    monkeypatch.setattr(documents, "MAX_EXTRACTED_CHARACTERS", 100, raising=False)
    with pytest.raises(KnowledgeBaseError) as error:
        parse(path, "text/plain")
    assert error.value.code == "document_too_complex"


@pytest.mark.asyncio
async def test_isolated_parser_round_trip_and_error(tmp_path: Path) -> None:
    function = getattr(documents, "parse_in_subprocess", None)
    assert function is not None, "尚未实现隔离解析"
    path = tmp_path / "资料.txt"
    path.write_text("行一\n行二😀", encoding="utf-8")
    blocks = await function(path, "text/plain")
    assert blocks[0].text == "行一\n行二😀" and blocks[0].source.line_end == 2
    with pytest.raises(KnowledgeBaseError) as error:
        await function(tmp_path / "missing.pdf", "application/pdf")
    assert error.value.code == "invalid_file"


@pytest.mark.asyncio
async def test_isolated_parser_accepts_maximum_size_plain_text(tmp_path: Path) -> None:
    function = getattr(documents, "parse_in_subprocess", None)
    assert function is not None, "尚未实现隔离解析"
    path = tmp_path / "large.txt"
    path.write_bytes(b"a" * 20_000_000)
    blocks = await function(path, "text/plain")
    assert len(blocks[0].text) == 20_000_000


@pytest.mark.skipif(os.environ.get("RAG_TOKENIZER_TESTS") != "1", reason="需要显式挂载已缓存的真实 Qwen 分词器")
def test_real_qwen_tokenizer_preserves_unique_multilingual_content(tmp_path: Path) -> None:
    from app.services.rag.clients import load_tokenizer
    tokenizer = load_tokenizer()
    path = tmp_path / "Qwen.md"
    original = "# Qwen验证\n\n" + "\n".join(f"第{index}条：中文😀 Text-{index}; 不得遗漏任何字符。" for index in range(100))
    path.write_text(original, encoding="utf-8")
    blocks = parse(path, "text/markdown")
    chunks = split(blocks, tokenizer, overlap=0)
    assert "".join("".join(item.text.split()) for item in chunks) == "".join(original.split())
    assert all(item.token_count == len(tokenizer.encode(item.text, add_special_tokens=True).ids) <= 128 for item in chunks)


@pytest.mark.skipif(os.environ.get("RAG_TOKENIZER_TESTS") != "1", reason="需要真实 Qwen 分词器缓存")
def test_token_valid_chunk_can_exceed_mysql_text_byte_limit(tmp_path: Path) -> None:
    from app.services.rag.clients import load_tokenizer
    tokenizer = load_tokenizer()
    path = tmp_path / "空白密集.txt"
    original = ("word" + " " * 128) * 600 + "end"
    path.write_text(original, encoding="utf-8")
    chunks = split(parse(path, "text/plain"), tokenizer, size=2048, overlap=0)
    assert len(chunks) == 1 and chunks[0].token_count <= 2048
    assert len(chunks[0].text.encode()) > 65_535
