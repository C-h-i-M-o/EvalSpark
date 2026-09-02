"""只在显式启动的隔离 Compose 中调用真实 TEI、Qdrant 和 Celery。"""

import asyncio
import io
import os
import time
from datetime import UTC, datetime, timedelta
from pathlib import Path
from uuid import uuid4

import pytest
from qdrant_client import AsyncQdrantClient
from sqlalchemy import select
from starlette.datastructures import Headers, UploadFile

from app.core.config import settings
from app.models.knowledge_base import KnowledgeBase, KnowledgeDocument, RagJob
from app.schemas.knowledge_base import KnowledgeBaseCreate, KnowledgeBasePatch
from app.services.knowledge_base_service import KnowledgeBaseService
from app.services.rag.clients import EmbeddingClient, VectorClient, VectorStore, VECTOR_COLLECTION
from app.services.rag.documents import MEDIA_TYPES

pytestmark = pytest.mark.skipif(os.environ.get("RAG_LIFECYCLE_TESTS") != "1", reason="需要独立完整 RAG 测试服务")


@pytest.mark.asyncio
async def test_real_qdrant_confirms_indexes_and_scoped_cleanup() -> None:
    assert settings.rag_qdrant_url.host == "qdrant-test"
    user_id = uuid4().int % 1_000_000_000 + 1_000_000
    client = AsyncQdrantClient(url=str(settings.rag_qdrant_url), check_compatibility=False, trust_env=False)
    store = VectorStore(client)
    try:
        await store.ensure_collection()
        await store.ensure_collection()
        point = str(uuid4())
        for _ in range(2):
            await store.upsert(user_id, 1, 1, 1, [(point, 0)], [[1.0] + [0.0] * 1023])
        assert await store.count(user_id, 1, 1) == 1
        info = await client.get_collection(VECTOR_COLLECTION)
        assert info.payload_schema["user_id"].params.is_principal is True
        matches = await VectorClient(client=client).search(user_id, 1, [(1, 1)], [1.0] + [0.0] * 1023)
        assert [match.chunk_id for match in matches] == [point]
        assert await VectorClient(client=client).search(user_id + 1, 1, [(1, 1)], [1.0] + [0.0] * 1023) == []
        await store.delete(user_id, 1, 1)
        assert await store.count(user_id, 1, 1) == 0
    finally:
        await store.delete(user_id, 1, 1)
        await client.close()


@pytest.mark.skipif(os.environ.get("RAG_RESTART_TESTS") != "1", reason="需要从宿主机先停止再启动隔离 Worker")
@pytest.mark.asyncio
async def test_real_worker_recovers_expired_cleanup_lease_without_embedding(rag_sessions, rag_users) -> None:
    from app.services.rag.documents import store_upload
    from app.services.rag.jobs import JobRepository
    assert settings.rag_documents_dir == Path("/test-documents")
    service = KnowledgeBaseService(publisher=lambda _: False)
    upload = UploadFile(io.BytesIO("需要清理的非敏感测试文件".encode()), filename="恢复测试.txt", headers=Headers({"content-type": "text/plain"}))
    stored = await store_upload(upload)
    async with rag_sessions() as db:
        kb = await service.create(db, rag_users[0], KnowledgeBaseCreate(name="清理重启验证"))
        # 模拟此前解析失败的文档，避免在模型关闭时创建无关索引任务。
        document = KnowledgeDocument(knowledge_base_id=kb.id, user_id=rag_users[0], **stored.__dict__, status="failed")
        db.add(document)
        await db.commit()
        cleanup = await service.delete(db, rag_users[0], kb.id)
    claim = await JobRepository(rag_sessions).claim(cleanup.job_id)
    assert claim is not None and claim.attempt == 1
    async with rag_sessions() as db:
        job = await db.get(RagJob, cleanup.job_id)
        job.lease_until = datetime.now(UTC).replace(tzinfo=None) - timedelta(minutes=15)
        job.stage = "deleting_vectors"
        await db.commit()
    print("RAG_RESTART_READY", flush=True)
    await wait_job(rag_sessions, cleanup.job_id)
    async with rag_sessions() as db:
        job = await db.get(RagJob, cleanup.job_id)
        assert job.attempt == 2 and job.status == "succeeded"
        assert (await db.get(KnowledgeBase, kb.id)).status == "deleted"
        assert (await db.get(KnowledgeDocument, document.id)).status == "deleted"
    assert not (settings.rag_documents_dir / stored.storage_key).exists()


async def wait_job(sessions, job_id: str, seconds: int = 180) -> None:
    deadline = time.monotonic() + seconds
    while time.monotonic() < deadline:
        async with sessions() as db:
            job = await db.get(RagJob, job_id)
            if job.status in ("succeeded", "failed"):
                assert job.status == "succeeded", (job.stage, job.error_code)
                return
        await asyncio.sleep(1)
    pytest.fail("真实 Worker 未在限定时间内完成作业")


@pytest.mark.asyncio
async def test_real_worker_indexes_four_formats_recovers_notification_and_cleans_up(rag_sessions, rag_users, tmp_path: Path) -> None:
    assert settings.rag_qdrant_url.host == "qdrant-test"
    assert settings.rag_embedding_url.host == "embedding-test"
    assert settings.rag_redis_url.host == "redis-test"
    assert settings.rag_documents_dir == Path("/test-documents")
    from docx import Document
    from test_rag_documents import make_pdf

    (tmp_path / "制度.txt").write_text("报销需要发票和主管审批。\n申请应在月底前提交。", encoding="utf-8")
    (tmp_path / "指南.md").write_text("# 请假指南\n\n请假需要提前三天提交申请，等待主管审批。", encoding="utf-8")
    make_pdf(tmp_path / "资料.pdf")
    word = Document()
    word.add_heading("培训", 1)
    word.add_table(rows=1, cols=1).cell(0, 0).text = "员工入职后参加培训。"
    word.save(tmp_path / "资料.docx")
    owner, other = rag_users
    service = KnowledgeBaseService()
    async with rag_sessions() as db:
        kb = await service.create(db, owner, KnowledgeBaseCreate(name="真实四格式验收"))
        jobs = []
        for path in sorted(tmp_path.iterdir()):
            upload = UploadFile(io.BytesIO(path.read_bytes()), filename=path.name, headers=Headers({"content-type": MEDIA_TYPES[path.suffix]}))
            # Markdown 故意漏发通知，验证恢复循环能从 MySQL 找回。
            uploader = KnowledgeBaseService(publisher=lambda _: False) if path.suffix == ".md" else service
            jobs.append(await uploader.upload(db, owner, kb.id, upload))
    for job in jobs:
        await wait_job(rag_sessions, job.job_id)
    client = AsyncQdrantClient(url=str(settings.rag_qdrant_url), check_compatibility=False, trust_env=False)
    try:
        info = await client.get_collection(VECTOR_COLLECTION)
        assert set(info.payload_schema) >= {"user_id", "knowledge_base_id", "document_id", "index_revision", "chunk_index"}
        assert info.payload_schema["user_id"].params.is_principal is True
        async with rag_sessions() as db:
            documents = (await db.scalars(select(KnowledgeDocument).where(KnowledgeDocument.knowledge_base_id == kb.id))).all()
            assert len(documents) == 4 and all(item.status == "ready" for item in documents)
            assert (await db.get(KnowledgeBase, kb.id)).status == "ready"
            versions = [(item.id, item.index_revision) for item in documents]
        query = (await EmbeddingClient().embed(["报销需要哪些材料？"], "query"))[0]
        matches = await VectorClient(client=client).search(owner, kb.id, versions, query)
        assert matches and {item.document_id for item in matches} <= {item.id for item in documents}
        assert await VectorClient(client=client).search(other, kb.id, versions, query) == []
        async with rag_sessions() as db:
            await service.patch(db, owner, kb.id, KnowledgeBasePatch(chunk_size=256))
            rebuilt = await service.reindex(db, owner, kb.id)
        await wait_job(rag_sessions, rebuilt.job_id)
        store = VectorStore(client)
        for document in documents:
            assert await store.count(owner, kb.id, document.id, 1) == 0
        async with rag_sessions() as db:
            deleted = await service.delete(db, owner, kb.id)
        await wait_job(rag_sessions, deleted.job_id)
        for document in documents:
            assert await store.count(owner, kb.id, document.id) == 0
            assert not (settings.rag_documents_dir / document.storage_key).exists()
        async with rag_sessions() as db:
            assert (await db.get(KnowledgeBase, kb.id)).status == "deleted"
    finally:
        await client.close()
